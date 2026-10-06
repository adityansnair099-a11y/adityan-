from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Theme colors
navy = RGBColor(11, 50, 88)
teal = RGBColor(33, 128, 146)
light = RGBColor(244, 247, 250)
text = RGBColor(34, 34, 34)
muted = RGBColor(96, 96, 96)
accent = RGBColor(224, 164, 58)


def add_title_slide(title, subtitle=''):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = light
    bar = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.35))
    bar.fill.solid()
    bar.fill.fore_color.rgb = navy
    bar.line.fill.background()

    title_box = slide.shapes.title
    title_box.text = title
    title_box.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
    title_box.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    title_box.text_frame.paragraphs[0].runs[0].font.bold = True
    title_box.text_frame.paragraphs[0].runs[0].font.color.rgb = navy

    if subtitle:
        subtitle_box = slide.placeholders[1]
        subtitle_box.text = subtitle
        subtitle_box.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
        subtitle_box.text_frame.paragraphs[0].runs[0].font.size = Pt(16)
        subtitle_box.text_frame.paragraphs[0].runs[0].font.color.rgb = muted

    return slide


def add_bullets_slide(title, bullets, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = light

    # Header accent bar
    bar = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.25))
    bar.fill.solid()
    bar.fill.fore_color.rgb = navy
    bar.line.fill.background()

    title_box = slide.shapes.title
    title_box.text = title
    title_box.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
    title_box.text_frame.paragraphs[0].runs[0].font.bold = True
    title_box.text_frame.paragraphs[0].runs[0].font.color.rgb = navy

    if subtitle:
        body = slide.shapes.placeholders[1]
        body.text = subtitle
        body.text_frame.paragraphs[0].font.size = Pt(18)
        body.text_frame.paragraphs[0].font.color.rgb = muted

    body_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.4), Inches(11.5), Inches(4.6))
    tf = body_box.text_frame
    tf.word_wrap = True
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.bullet = True
        p.space_after = Pt(8)
        p.font.size = Pt(20)
        p.font.color.rgb = text

    return slide


def add_section_slide(title, left_heading, left_points, right_heading, right_points):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = light

    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.3), Inches(12), Inches(0.6))
    title_frame = title_box.text_frame
    p = title_frame.paragraphs[0]
    p.text = title
    p.alignment = PP_ALIGN.LEFT
    p.font.bold = True
    p.font.size = Pt(22)
    p.font.color.rgb = navy

    # left block
    left = slide.shapes.add_shape(1, Inches(0.7), Inches(1.0), Inches(5.4), Inches(4.8))
    left.fill.solid()
    left.fill.fore_color.rgb = RGBColor(255, 255, 255)
    left.line.color.rgb = navy
    left.line.width = Pt(1.2)
    tf_left = left.textframe
    p = tf_left.paragraphs[0]
    p.text = left_heading
    p.level = 0
    p.font.bold = True
    p.font.size = Pt(20)
    p.font.color.rgb = navy
    for bullet in left_points:
        p = tf_left.add_paragraph()
        p.text = bullet
        p.bullet = True
        p.level = 0
        p.font.size = Pt(17)
        p.font.color.rgb = text
        p.space_after = Pt(8)

    # right block
    right = slide.shapes.add_shape(1, Inches(6.8), Inches(1.0), Inches(5.5), Inches(4.8))
    right.fill.solid()
    right.fill.fore_color.rgb = RGBColor(255, 255, 255)
    right.line.color.rgb = teal
    right.line.width = Pt(1.2)
    tf_right = right.textframe
    p = tf_right.paragraphs[0]
    p.text = right_heading
    p.font.bold = True
    p.font.size = Pt(20)
    p.font.color.rgb = teal
    for bullet in right_points:
        p = tf_right.add_paragraph()
        p.text = bullet
        p.bullet = True
        p.level = 0
        p.font.size = Pt(17)
        p.font.color.rgb = text
        p.space_after = Pt(8)

    return slide


# Slide 1
add_title_slide(
    'SkyCity Auckland Restaurants & Bars',
    'Research Paper | Exploratory Data Analysis, Insights, and Strategic Recommendations'
)

# Slide 2
add_bullets_slide(
    '1. Research Objective',
    [
        'Assess how restaurant format, geography, and channel mix affect commercial performance.',
        'Examine the relationship between delivery dependence, operating costs, and net profitability.',
        'Translate the findings into practical strategic recommendations for portfolio optimization.'
    ]
)

# Slide 3
add_section_slide(
    '2. Dataset and Business Context',
    'Data included',
    [
        'Cuisine type, restaurant format, subregion, and growth factor',
        'Monthly orders, average order value, revenue by channel',
        'COGS, OPEX, commission rates, and delivery cost per order',
        'Net profit by channel: in-store, UberEats, DoorDash, and self-delivery'
    ],
    'Why it matters',
    [
        'Revenue alone is insufficient without contribution margin analysis',
        'Delivery channels increase reach but add cost and margin pressure',
        'Portfolio performance depends on unit economics, not just top-line sales',
        'Regional and format differences matter for strategy design'
    ]
)

# Slide 4
add_bullets_slide(
    '3. Exploratory Findings',
    [
        'Digital growth is a double-edged sword: it expands demand but increases commission and delivery costs.',
        'Restaurant format strongly influences economics; QSR and café operations tend to favor repeat ordering and efficient volume.',
        'Location and subregion shape customer behavior, affecting demand mix, service expectations, and cost-to-serve.',
        'The strongest units are those that combine attractive AOV, manageable costs, and disciplined channel mix.'
    ]
)

# Slide 5
add_section_slide(
    '4. Key Insights',
    'Main performance patterns',
    [
        'High delivery dependence can weaken profit despite strong revenue growth',
        'Not all high-volume restaurants are highly profitable',
        'Strong in-store economics remain the most resilient foundation',
        'Operational efficiency matters as much as sales expansion'
    ],
    'Strategic implication',
    [
        'The portfolio should be managed by contribution margin and channel health, not just revenue size',
        'Low-margin growth should be challenged with pricing, menu, or delivery redesign',
        'The best-performing stores balance accessibility with cost discipline'
    ]
)

# Slide 6
add_bullets_slide(
    '5. Recommendations',
    [
        'Protect and expand high-margin in-store channels before over-investing in delivery-heavy growth.',
        'Use delivery strategically, especially where it supports acquisition and convenience without erosion of margin.',
        'Review underperforming restaurants by format, subregion, and commission burden before scaling investment.',
        'Improve menu design and pricing architecture to sustain AOV while controlling cost-to-serve.'
    ]
)

# Slide 7
add_bullets_slide(
    '6. Conclusion',
    [
        'The restaurant portfolio performs best when commercial growth is matched by unit-level profitability discipline.',
        'The key question is not simply how much a restaurant sells, but how much it retains after commissions, delivery costs, and operating expenses.',
        'A data-led portfolio strategy can improve profitability by reallocating investment toward formats and locations with stronger economic resilience.'
    ]
)

# Save presentation
output_path = r"c:\Users\DELL\OneDrive\Desktop\sky blue\SkyCity_Auckland_Restaurants_Research_Presentation.pptx"
prs.save(output_path)
print(output_path)
