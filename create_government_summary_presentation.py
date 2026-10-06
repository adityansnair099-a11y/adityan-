from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


BASE = Path(__file__).parent
OUTPUT = BASE / "SkyCity_Auckland_Government_Executive_Summary.pptx"

NAVY = RGBColor(18, 59, 95)
TEAL = RGBColor(45, 136, 149)
GOLD = RGBColor(217, 164, 65)
LIGHT = RGBColor(244, 247, 250)
WHITE = RGBColor(255, 255, 255)
TEXT = RGBColor(31, 42, 55)
MUTED = RGBColor(95, 104, 115)
RED = RGBColor(176, 73, 63)
LINE = RGBColor(220, 228, 235)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def add_text(slide, text, x, y, width, height, size=18, color=TEXT, bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.02)
    frame.margin_right = Inches(0.02)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    paragraph.font.name = "Aptos"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    return box


def add_base(title, page):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = LIGHT
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.18))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    add_text(slide, title, 0.65, 0.42, 12.0, 0.55, 25, NAVY, True)
    add_text(slide, f"SKYCITY AUCKLAND  |  GOVERNMENT SUMMARY  |  {page}",
             0.65, 7.12, 12.0, 0.18, 8, MUTED)
    return slide


def add_metric(slide, x, y, width, value, label, color=TEAL):
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x), Inches(y), Inches(width), Inches(1.25),
    )
    card.fill.solid()
    card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = LINE
    card.line.width = Pt(0.8)
    add_text(slide, value, x + 0.18, y + 0.16, width - 0.36, 0.48, 25, color, True)
    add_text(slide, label, x + 0.18, y + 0.70, width - 0.36, 0.34, 11, MUTED)


def add_bullet(slide, text, x, y, width, color=TEXT, size=16):
    add_text(slide, "•", x, y, 0.25, 0.42, size, TEAL, True)
    add_text(slide, text, x + 0.28, y, width - 0.28, 0.68, size, color)


def add_segment_bar(slide, name, margin, y):
    baseline = 4.65
    scale = 0.23
    add_text(slide, name, 0.95, y - 0.05, 2.8, 0.35, 15, TEXT, True)
    marker = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(baseline), Inches(y), Inches(0.025), Inches(0.28)
    )
    marker.fill.solid()
    marker.fill.fore_color.rgb = MUTED
    marker.line.fill.background()
    width = abs(margin) * scale
    x = baseline - width if margin < 0 else baseline
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y + 0.025), Inches(width), Inches(0.23)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = RED if margin < 0 else TEAL
    bar.line.fill.background()
    label_x = baseline + 0.12 if margin < 0 else baseline + width + 0.12
    add_text(slide, f"{margin:.1f}%", label_x, y - 0.04, 0.85, 0.34, 14,
             RED if margin < 0 else NAVY, True)


slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = LIGHT
top = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.22))
top.fill.solid()
top.fill.fore_color.rgb = NAVY
top.line.fill.background()
add_text(slide, "SKYCITY AUCKLAND", 0.8, 1.1, 11.8, 0.35, 14, TEAL, True)
add_text(slide, "Hospitality portfolio", 0.8, 1.62, 11.7, 0.78, 34, NAVY, True)
add_text(slide, "Executive summary for government stakeholders", 0.8, 2.48,
         11.6, 0.45, 20, MUTED)
add_metric(slide, 0.8, 3.65, 2.8, "$77.74m", "Recorded revenue")
add_metric(slide, 3.8, 3.65, 2.8, "$7.87m", "Recorded net profit")
add_metric(slide, 6.8, 3.65, 2.8, "10.1%", "Aggregate margin", NAVY)
add_metric(slide, 9.8, 3.65, 2.8, "24.1%", "Outlet records in loss", RED)
add_text(slide, "Portfolio-level diagnostic based on the supplied outlet dataset.",
         0.82, 5.45, 11.5, 0.38, 14, TEXT)
add_text(slide, "Currency and reporting period are not specified in the source file.",
         0.82, 5.9, 11.5, 0.35, 11, MUTED)
add_text(slide, "SKYCITY AUCKLAND  |  GOVERNMENT SUMMARY  |  1", 0.82, 7.12,
         11.7, 0.18, 8, MUTED)


slide = add_base("Overall performance hides format differences", 2)
add_metric(slide, 0.8, 1.22, 3.0, "1,696", "Outlet records analyzed", NAVY)
add_metric(slide, 4.0, 1.22, 3.0, "1,288", "Non-loss-making records", TEAL)
add_metric(slide, 7.2, 1.22, 3.0, "408", "Loss-making records", RED)
add_text(slide, "Aggregate segment net margin", 0.95, 3.0, 4.0, 0.35, 17, NAVY, True)
add_segment_bar(slide, "Ghost Kitchen", 22.2, 3.65)
add_segment_bar(slide, "Cafe", 16.2, 4.35)
add_segment_bar(slide, "QSR", 15.1, 5.05)
add_segment_bar(slide, "Full-service", -5.8, 5.75)
add_text(slide, "Margins use total segment net profit divided by total segment revenue.",
         0.95, 6.43, 11.4, 0.32, 11, MUTED)


slide = add_base("Channel results are uneven", 3)
add_text(slide, "Recorded net profit by channel", 0.85, 1.2, 5.3, 0.35, 17, NAVY, True)
channel_values = [
    ("In-store", 3.828, NAVY),
    ("Self-delivery", 3.628, TEAL),
    ("Uber Eats", 0.258, GOLD),
    ("DoorDash", 0.151, RED),
]
for i, (label, amount, color) in enumerate(channel_values):
    y = 1.9 + i * 0.82
    add_text(slide, label, 0.9, y, 1.7, 0.32, 14, TEXT, True)
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(2.65), Inches(y + 0.025),
        Inches(max(0.2, amount * 1.03)), Inches(0.25),
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()
    add_text(slide, f"${amount:.3f}m", 6.8, y - 0.03, 1.65, 0.32, 14, TEXT, True)
add_text(slide, "Average order share across outlet records", 8.75, 1.2,
         3.8, 0.35, 16, NAVY, True)
for i, (label, value) in enumerate([
    ("In-store", "17.5%"), ("Uber Eats", "40.3%"),
    ("DoorDash", "22.0%"), ("Self-delivery", "20.1%"),
]):
    y = 1.92 + i * 0.63
    add_text(slide, label, 8.8, y, 2.0, 0.3, 13, TEXT)
    add_text(slide, value, 11.1, y, 1.15, 0.3, 13, NAVY, True, PP_ALIGN.RIGHT)
add_metric(slide, 8.85, 4.9, 3.3, "30.0%", "Average commission rate", GOLD)
add_text(slide, "Aggregate channel profits are descriptive; they do not show that channel mix causes profitability differences.",
         0.9, 6.25, 11.7, 0.55, 12, MUTED)


slide = add_base("Implications for public decision-makers", 4)
add_text(slide, "What the data support", 0.85, 1.2, 5.4, 0.4, 18, NAVY, True)
add_bullet(slide, "Assess businesses by format and outlet-level economics; portfolio averages conceal losses.",
           0.9, 1.9, 5.45)
add_bullet(slide, "Use channel-level costs and commissions as indicators for further investigation, not proof of cause.",
           0.9, 2.95, 5.45)
add_bullet(slide, "Consider targeted diagnostics and cost-management support where local evidence identifies need.",
           0.9, 4.0, 5.45)
add_text(slide, "Evidence needed before policy or funding decisions", 6.85, 1.2,
         5.5, 0.58, 17, TEAL, True)
add_bullet(slide, "Employment, wages, business survival, and local procurement.",
           6.9, 1.9, 5.45)
add_bullet(slide, "Verified source, reporting period, currency, and profit-calculation method.",
           6.9, 2.95, 5.45)
add_bullet(slide, "Comparable trend data and neighborhood-level outcomes.",
           6.9, 4.0, 5.45)
add_text(slide, "The dataset does not measure wider employment, tourism, tax, or community impacts.",
         0.9, 5.65, 11.6, 0.55, 13, MUTED)


slide = add_base("Data, method, and limitations", 5)
add_bullet(slide, "Scope: 1,696 outlet records in the supplied restaurants and bars CSV.",
           0.9, 1.35, 11.5, size=15)
add_bullet(slide, "Total revenue sums the four channel revenue fields; total net profit sums the four channel net-profit fields.",
           0.9, 2.25, 11.5, size=15)
add_bullet(slide, "Segment margins are calculated from aggregate segment profit and revenue, not as an average of outlet margins.",
           0.9, 3.15, 11.5, size=15)
add_bullet(slide, "The source does not identify currency, reporting period, provenance, or calculation methodology; figures are not independently audited.",
           0.9, 4.05, 11.5, size=15)
add_bullet(slide, "Treat these results as a portfolio diagnostic, not verified sector-wide statistics or causal evidence.",
           0.9, 4.95, 11.5, size=15)


prs.save(OUTPUT)
print(OUTPUT)