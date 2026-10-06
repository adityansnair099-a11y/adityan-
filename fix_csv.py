import csv
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

path = Path(__file__).with_name("SkyCity Auckland Restaurants & Bars.csv")
output_path = path.with_suffix(".xlsx")
with path.open("r", newline="", encoding="utf-8-sig") as f:
    reader = csv.DictReader(f)
    rows = list(reader)

if not rows:
    raise SystemExit("CSV is empty")

fieldnames = reader.fieldnames
share_columns = ["InStoreShare", "UE_share", "DD_share", "SD_share"]
for row in rows:
    total = float(row["MonthlyOrders"] or 0)
    channel_orders = [
        float(row["InStoreOrders"] or 0),
        float(row["UberEatsOrders"] or 0),
        float(row["DoorDashOrders"] or 0),
        float(row["SelfDeliveryOrders"] or 0),
    ]
    if total == 0:
        shares = [0.0, 0.0, 0.0, 0.0]
    else:
        shares = [orders / total for orders in channel_orders]
    for column, share in zip(share_columns, shares):
        row[column] = share

workbook = Workbook()
worksheet = workbook.active
worksheet.title = "Restaurants & Bars"
worksheet.append(fieldnames)

for row in rows:
    values = []
    for column in fieldnames:
        value = row[column]
        if column in share_columns:
            values.append(value)
            continue
        try:
            number = float(value)
            values.append(int(number) if number.is_integer() else number)
        except (TypeError, ValueError):
            values.append(value)
    worksheet.append(values)

worksheet.freeze_panes = "A2"
worksheet.sheet_view.showGridLines = False
for cell in worksheet[1]:
    cell.font = Font(color="FFFFFF", bold=True)
    cell.fill = PatternFill("solid", fgColor="17365D")

table = Table(displayName="RestaurantData", ref=worksheet.dimensions)
table.tableStyleInfo = TableStyleInfo(
    name="TableStyleMedium2",
    showFirstColumn=False,
    showLastColumn=False,
    showRowStripes=True,
    showColumnStripes=False,
)
worksheet.add_table(table)

currency_columns = {
    "AOV", "InStoreRevenue", "UberEatsRevenue", "DoorDashRevenue",
    "SelfDeliveryRevenue", "DeliveryCostPerOrder", "SD_DeliveryTotalCost",
    "InStoreNetProfit", "UberEatsNetProfit", "DoorDashNetProfit",
    "SelfDeliveryNetProfit",
}
integer_columns = {
    "RestaurantID", "MonthlyOrders", "InStoreOrders", "UberEatsOrders",
    "DoorDashOrders", "SelfDeliveryOrders",
}
percentage_columns = set(share_columns) | {"COGSRate", "OPEXRate", "CommissionRate"}
for column_index, column in enumerate(fieldnames, 1):
    letter = worksheet.cell(row=1, column=column_index).column_letter
    worksheet.column_dimensions[letter].width = min(max(len(column) + 3, 14), 30)
    for cells in worksheet.iter_cols(min_col=column_index, max_col=column_index, min_row=2):
        for value_cell in cells:
            if column in percentage_columns:
                value_cell.number_format = "0.0%"
            elif column in currency_columns:
                value_cell.number_format = '$#,##0.00;[Red]-$#,##0.00'
            elif column in integer_columns:
                value_cell.number_format = "#,##0"
            elif isinstance(value_cell.value, (int, float)):
                value_cell.number_format = "#,##0.00"

workbook.save(output_path)
check = load_workbook(output_path, read_only=True, data_only=True)
print(f"created {output_path.name} with {worksheet.max_row - 1} rows")
print(f"sample share: {check.active['AA2'].value:.1%}")
